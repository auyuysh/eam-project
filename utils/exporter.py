import io
import logging
from datetime import datetime

from flask import send_file
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer,
)

logger = logging.getLogger("eam.exporter")

_PDF_PAGE_WIDTH = landscape(A4)[0]
_TIMESTAMP = lambda: datetime.now().strftime("%Y%m%d_%H%M%S")


def generate_generic_excel(filename, column_headers, data_rows):
    """
    Build a styled .xlsx workbook and stream it as an attachment.

    Parameters
    ----------
    filename : str
        Download name, e.g. ``'it_assets_list.xlsx'``.
    column_headers : list[str]
        Human-readable header labels.
    data_rows : list[tuple | list]
        Raw cell values in column order matching *column_headers*.

    Returns
    -------
    flask.Response
        A ``send_file`` response with the spreadsheet payload.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Exported Data"

    header_font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="1A1A2E", end_color="1A1A2E", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_font = Font(name="Calibri", size=11)
    thin_border = Border(
        left=Side(style="thin", color="D0D0D0"),
        right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"),
        bottom=Side(style="thin", color="D0D0D0"),
    )

    for col_idx, header in enumerate(column_headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    for row_idx, row_data in enumerate(data_rows, 2):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = cell_font
            cell.alignment = Alignment(vertical="center")
            cell.border = thin_border
            if row_idx % 2 == 0:
                cell.fill = PatternFill(
                    start_color="F9FAFB", end_color="F9FAFB", fill_type="solid"
                )

    for col_idx in range(1, len(column_headers) + 1):
        max_len = len(str(column_headers[col_idx - 1]))
        for row_idx in range(2, len(data_rows) + 2):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[
            ws.cell(row=1, column=col_idx).column_letter
        ].width = min(max_len + 4, 50)

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)

    logger.info(
        "Excel export generated: %s (%d rows)", filename, len(data_rows)
    )

    return send_file(
        file_stream,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )


def generate_generic_pdf(filename, title, column_headers, data_rows):
    """
    Build a landscape A4 PDF table and stream it as an attachment.

    Parameters
    ----------
    filename : str
        Download name, e.g. ``'it_assets_report.pdf'``.
    title : str
        Document heading rendered at the top of the page.
    column_headers : list[str]
        Table column labels.
    data_rows : list[tuple | list]
        Raw cell values in column order matching *column_headers*.

    Returns
    -------
    flask.Response
        A ``send_file`` response with the PDF payload.
    """
    file_stream = io.BytesIO()
    doc = SimpleDocTemplate(
        file_stream,
        pagesize=landscape(A4),
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ExportTitle",
        parent=styles["Heading1"],
        fontSize=16,
        textColor=colors.HexColor("#1A1A2E"),
        spaceAfter=4 * mm,
    )
    meta_style = ParagraphStyle(
        "ExportMeta",
        parent=styles["Normal"],
        fontSize=9,
        textColor=colors.grey,
        spaceAfter=8 * mm,
    )
    cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontSize=8,
        leading=10,
    )
    header_cell_style = ParagraphStyle(
        "TableHeaderCell",
        parent=styles["Normal"],
        fontSize=8,
        leading=10,
        textColor=colors.white,
    )

    elements = []
    elements.append(Paragraph(str(title), title_style))
    elements.append(
        Paragraph(
            f"Generated on {datetime.now().strftime('%d %b %Y, %I:%M %p')}",
            meta_style,
        )
    )

    table_data = [
        [Paragraph(str(h), header_cell_style) for h in column_headers]
    ]
    for row in data_rows:
        table_data.append(
            [Paragraph(str(cell) if cell is not None else "", cell_style) for cell in row]
        )

    if not data_rows:
        table_data.append(
            [Paragraph("(No records found)", cell_style)] + ["" for _ in column_headers[1:]]
        )

    avail_width = _PDF_PAGE_WIDTH - 30 * mm
    col_count = len(column_headers)
    col_width = avail_width / col_count

    tbl = Table(table_data, colWidths=[col_width] * col_count, repeatRows=1)
    tbl.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1A1A2E")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, 0), 8),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
                ("TOPPADDING", (0, 0), (-1, 0), 6),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D0D0D0")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 1), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 1), (-1, -1), 4),
            ]
            + [
                (
                    "BACKGROUND",
                    (0, r),
                    (-1, r),
                    colors.HexColor("#F9FAFB"),
                )
                for r in range(2, len(table_data), 2)
            ]
        )
    )
    elements.append(tbl)

    doc.build(elements)
    file_stream.seek(0)

    logger.info(
        "PDF export generated: %s (%d rows)", filename, len(data_rows)
    )

    return send_file(
        file_stream,
        mimetype="application/pdf",
        as_attachment=True,
        download_name=filename,
    )
